"""One batch, one output folder. The same validated MSG builder serves app and CLI."""
import csv
import fcntl
import hashlib
import io
import json
import os
import re
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import load_workbook

from .config import atomic_write, path_for, write_json
from .drafts import _filename_part, build_draft, draft_filename, file_hash, prepare_draft


def _digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _list_rows(path, sheet=None):
    if path.suffix.lower() == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            ws = wb[sheet] if sheet else wb.active
            rows = []
            for number, cells in enumerate(ws.iter_rows(), 1):
                if any(cell.data_type in ("f", "e") for cell in cells):
                    raise ValueError("Liste, Zeile %d: Formeln/Excel-Fehler zuerst durch geprüfte Werte ersetzen." % number)
                rows.append([cell.value for cell in cells])
            return rows, ws.title
        finally:
            wb.close()
    text = path.read_text(encoding="utf-8-sig")
    if not text.strip():
        return [], None
    try:
        delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ";" if ";" in text.splitlines()[0] else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)), None


def import_list(cfg, source, sheet=None):
    """A row already contains its individual subject/body; no hidden text generation."""
    source = Path(source).expanduser().resolve()
    try:
        rows, sheet_name = _list_rows(source, sheet)
    except (csv.Error, UnicodeError, zipfile.BadZipFile, KeyError) as error:
        raise ValueError("Liste konnte nicht gelesen werden: " + str(error)) from error
    if not rows:
        raise ValueError("Die Liste ist leer.")
    aliases = {"e-mail": "to", "email": "to", "an": "to", "empfänger": "to",
               "betreff": "subject", "text": "body", "nachricht": "body",
               "name": "name", "anhänge": "attachments", "anhaenge": "attachments",
               "anhang": "attachments", "deck": "include_deck"}
    headers = [aliases.get(str(x or "").strip().lower(), str(x or "").strip().lower()) for x in rows[0]]
    known = {"to", "cc", "bcc", "subject", "body", "name", "attachments", "include_deck", "importance"}
    if not {"to", "subject", "body"}.issubset(headers):
        raise ValueError("Die Liste braucht die Spalten E-Mail, Betreff und Text. Eine reine Kontaktliste zuerst mit der KI um individuelle Texte ergänzen.")
    if any(headers.count(key) > 1 for key in known):
        raise ValueError("Die Liste enthält dieselbe Mail-Spalte mehrfach.")
    drafts = []
    for line, row in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if len(row) > len(headers) and any(str(value or "").strip() for value in row[len(headers):]):
            raise ValueError("Liste, Zeile %d: Mehr Werte als Spalten; CSV-Trennzeichen und Anführungszeichen prüfen." % line)
        data = {key: str(value).strip() if value is not None else "" for key, value in zip(headers, row) if key in known}
        meta = {key: data.get(key, "") for key in ("subject", "to", "cc", "bcc")}
        if data.get("name") and re.fullmatch(r'[^\s@<>,;]+@[^\s@<>,;]+', meta["to"]):
            name = data["name"].replace("\\", "\\\\").replace('"', '\\"')
            meta["to"] = '"' + name + '" <' + meta["to"] + '>'
        flag = data.get("include_deck", "").lower()
        if flag not in ("", "false", "nein", "0", "true", "ja", "1", "x"):
            raise ValueError("Liste, Zeile %d: Deck muss ja/nein sein." % line)
        meta["include_deck"] = flag in ("true", "ja", "1", "x")
        value = data.get("attachments", "")
        try:
            # Several attachments: "a.pdf; b.pdf" (or a JSON list, as before).
            files = json.loads(value) if value.startswith("[") else [item.strip() for item in value.split(";") if item.strip()]
        except json.JSONDecodeError as error:
            raise ValueError("Liste, Zeile %d: Mehrere Anhänge mit Semikolon trennen, z. B. a.pdf; b.pdf." % line) from error
        if not isinstance(files, list) or not all(isinstance(item, str) and item for item in files):
            raise ValueError("Liste, Zeile %d: Ungültige Anhänge." % line)
        meta["attachments"] = [str((source.parent / Path(item).expanduser()).resolve()) for item in files]
        if data.get("importance"):
            meta["importance"] = data["importance"]
        drafts.append((line, data.get("name") or data.get("to") or "Mail", meta, data.get("body", "")))
    if not drafts:
        raise ValueError("Die Liste enthält keine Empfängerzeilen.")
    base = cfg["root"] / "Entwürfe"
    base.mkdir(parents=True, exist_ok=True)
    folder = Path(tempfile.mkdtemp(prefix=_filename_part(source.stem, 45) + "-", dir=base))
    paths = []
    provenance = []
    for number, (line, name, meta, body) in enumerate(drafts, 1):
        path = folder / ("%04d – %s.md" % (number, _filename_part(name, 70)))
        atomic_write(path, "---\n" + yaml.safe_dump(meta, allow_unicode=True, sort_keys=False) + "---\n" + body + "\n")
        paths.append(path)
        provenance.append({"row": line, "markdown": path.name})
    write_json(folder / ".liste.json", {"source": str(source), "sha256": file_hash(source), "sheet": sheet_name, "rows": provenance})
    return paths, folder


def _sources(cfg, sources, sheet):
    paths = [Path(value).expanduser().resolve() for value in sources]
    if len(paths) == 1 and paths[0].suffix.lower() in (".csv", ".xlsx") and paths[0].is_file():
        return import_list(cfg, paths[0], sheet)
    if len(paths) == 1 and paths[0].is_dir():
        paths = sorted((p for p in paths[0].iterdir() if p.is_file() and p.suffix.lower() in (".md", ".markdown") and not p.name.startswith(".")), key=lambda p: p.name.casefold())
    if not paths:
        raise ValueError("Keine Markdown-Entwürfe im ausgewählten Ordner.")
    if any(not p.is_file() or p.suffix.lower() not in (".md", ".markdown") for p in paths):
        raise ValueError("Wähle Markdown-Dateien, einen Markdown-Ordner oder eine CSV-/XLSX-Liste.")
    if len(set(paths)) != len(paths):
        raise ValueError("Derselbe Markdown-Entwurf wurde mehrfach ausgewählt.")
    return paths, paths[0].parent


def _summary(state, folder):
    entries = state["entries"]
    created = sum(item["status"] == "fertig" for item in entries)
    errors = [{"number": item["number"], "source": item["source"], "error": item["error"]} for item in entries if item.get("error")]
    # Allowed, but often a leftover second draft for the same person: report, don't block.
    to = {}
    for item in entries:
        for address, _ in item.get("recipients", {}).get("to", []):
            to.setdefault(address.casefold(), []).append(item["number"])
    return {"batch": True, "folder": str(folder), "drafts_folder": state["drafts_folder"],
            "overview": str(folder / "Übersicht.csv"), "total": len(entries), "created": created,
            "failed": len(errors), "complete": created == len(entries) and not errors,
            "phase": state["phase"], "errors": errors,
            "repeated_to": [{"address": address, "numbers": numbers} for address, numbers in to.items() if len(numbers) > 1]}


def _persist(state, folder, progress_file):
    write_json(folder / ".auftrag.json", state)
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(["Nr", "An", "Cc", "Bcc", "Betreff", "Status", "MSG-Datei", "Markdown", "Fehler"])
    for item in state["entries"]:
        recipients = item.get("recipients", {})
        row = [item["number"], *["; ".join((name + " <" + address + ">") if name else address for address, name in recipients.get(key, [])) for key in ("to", "cc", "bcc")],
               item.get("subject", ""), item["status"], item.get("filename", ""), item["source"], item.get("error", "")]
        # Excel must display subjects/paths as text, not interpret them as formulas.
        writer.writerow(["'" + str(value) if str(value).startswith(("=", "+", "-", "@", "\t", "\r")) else value for value in row])
    atomic_write(folder / "Übersicht.csv", "\ufeff" + buffer.getvalue())
    if progress_file:
        write_json(progress_file, _summary(state, folder))


def _fingerprint(prepared):
    return _digest({"markdown": prepared["source_sha256"], "attachments": prepared["attachments"],
                    "images": [{k: v for k, v in image.items() if k != "path"} for image in prepared.get("images", [])],
                    "signature_html": prepared["signature_html"], "signature_text": prepared["signature_text"]})


def create_batch(cfg, sources=(), name=None, output=None, resume=None, progress_file=None, sheet=None):
    if bool(sources) == bool(resume):
        raise ValueError("Entweder Quellen oder --resume mit einem bisherigen Ausgabeordner angeben.")
    if resume:
        folder = Path(resume).expanduser().resolve()
        if not (folder / ".auftrag.json").is_file():
            raise ValueError("Dieser Ordner enthält keinen fortsetzbaren Sammel-Export.")
    else:
        paths, drafts_folder = _sources(cfg, sources, sheet)
        root = Path(output).expanduser().resolve() if output else path_for(cfg, "output_dir")
        root.mkdir(parents=True, exist_ok=True)
        label = name or (Path(sources[0]).stem if len(sources) == 1 else "E-Mails")
        folder = Path(tempfile.mkdtemp(prefix=datetime.now().strftime("%Y%m%d-%H%M%S-") + _filename_part(label, 45) + "-", dir=root))
    with (folder / ".lock").open("a+b") as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as error:
            raise ValueError("Dieser Sammel-Export läuft bereits. Bitte dessen Abschluss abwarten.") from error
        if resume:
            state = json.loads((folder / ".auftrag.json").read_text(encoding="utf-8"))
            if not isinstance(state, dict) or state.get("version") != 1 or not isinstance(state.get("entries"), list) or not state["entries"] or not isinstance(state.get("drafts_folder"), str):
                raise ValueError("Unbekanntes oder beschädigtes Auftragsformat.")
            for number, item in enumerate(state["entries"], 1):
                if not isinstance(item, dict) or not isinstance(item.get("source"), str) or not isinstance(item.get("status"), str):
                    raise ValueError("Beschädigter Eintrag im Auftragsformat.")
                filename = item.get("filename", "")
                if item.get("number") != number or (filename and (Path(filename).name != filename or not filename.endswith(".msg"))):
                    raise ValueError("Beschädigtes Auftragsformat: Dateinamen/Nummern stimmen nicht.")
                if item.get("msg_sha256") and not all(isinstance(item.get(key), str) and re.fullmatch(r"[0-9a-f]{64}", item[key]) for key in ("msg_sha256", "fingerprint")):
                    raise ValueError("Beschädigte Prüfsummen im Auftragsformat.")
        else:
            state = {"version": 1, "drafts_folder": str(drafts_folder), "phase": "prüfe",
                     "entries": [{"number": i, "source": str(path), "status": "wartet"} for i, path in enumerate(paths, 1)]}
        try:
            state["phase"] = "prüfe"
            _persist(state, folder, progress_file)
            prepared = {}
            duplicates = {}
            for item in state["entries"]:
                item.pop("error", None)
                try:
                    value = prepare_draft(cfg, item["source"])
                    if not value["recipients"]["to"]:
                        raise ValueError("An-Empfänger fehlt.")
                    fingerprint = _fingerprint(value)
                    filename = "%04d – %s.msg" % (item["number"], draft_filename(value["metadata"]["subject"], value["recipients"]))
                    if item.get("msg_sha256") and (item["fingerprint"] != fingerprint or item["filename"] != filename):
                        raise ValueError("Eingaben einer bereits erstellten Mail wurden geändert. Dafür einen neuen Stapel erstellen; die bisherige MSG bleibt erhalten.")
                    item.update(subject=value["metadata"]["subject"], recipients=value["recipients"], filename=filename, fingerprint=fingerprint)
                    target = folder / filename
                    if target.exists():
                        if not item.get("msg_sha256") or file_hash(target) != item["msg_sha256"]:
                            raise ValueError("Vorhandene MSG wurde außerhalb des Exports geändert; sie wird nicht überschrieben.")
                        item["status"] = "fertig"
                    else:
                        item["status"] = "wartet"
                    key = _digest([value["metadata"]["subject"], value["body"], value["recipients"], value["attachments"]])
                    duplicates.setdefault(key, []).append(item)
                    prepared[item["number"]] = value
                except Exception as error:
                    item.update(status="fehler", error=str(error) or type(error).__name__)
            for matches in duplicates.values():
                if len(matches) > 1:
                    numbers = ", ".join(str(item["number"]) for item in matches)
                    for item in matches:
                        item.update(status="fehler", error="Identischer Entwurf mehrfach vorhanden (Nr. " + numbers + "). Dubletten entfernen und neuen Stapel erstellen.")
            if any(item.get("error") for item in state["entries"]):
                state["phase"] = "Prüfung fehlgeschlagen"
                _persist(state, folder, progress_file)
                return _summary(state, folder)
            state["phase"] = "exportiere"
            _persist(state, folder, progress_file)
            for item in state["entries"]:
                if item["status"] == "fertig":
                    continue
                item["status"] = "erstelle"
                _persist(state, folder, progress_file)
                try:
                    target = folder / item["filename"]
                    pending = folder / (".%04d.pending.msg" % item["number"])
                    if not (item.get("msg_sha256") and pending.exists() and file_hash(pending) == item["msg_sha256"]):
                        data, html = build_draft(prepared[item["number"]])
                        atomic_write(folder / ".vorschau" / (item["filename"][:-4] + ".html"), html)
                        atomic_write(pending, data)
                        item["msg_sha256"] = hashlib.sha256(data).hexdigest()
                        _persist(state, folder, progress_file)
                    if target.exists():
                        raise ValueError("Zieldatei wurde während des Exports angelegt; sie wird nicht überschrieben.")
                    # Hard-link publication is atomic and refuses an existing target.
                    if hasattr(os, "link"):
                        os.link(pending, target)
                        pending.unlink()
                    else:  # Browser (Pyodide): single process, no hard links.
                        os.replace(pending, target)
                    item["status"] = "fertig"
                except Exception as error:
                    item.update(status="fehler", error=str(error) or type(error).__name__)
                _persist(state, folder, progress_file)
            state["phase"] = "fertig" if all(item["status"] == "fertig" for item in state["entries"]) else "unvollständig"
        except KeyboardInterrupt:
            state["phase"] = "unterbrochen"
        _persist(state, folder, progress_file)
        return _summary(state, folder)
